from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from excel_cleaner.cli import main
from excel_cleaner.pipeline import Options, clean, read_table, summarize, write_report

SAMPLE = Path(__file__).parent.parent / "sample-data" / "customers_messy.csv"
OPTS = Options(
    key_columns=["Order ID"],
    date_columns=["Order Date"],
    number_columns=["Amount"],
    email_columns=["Email"],
    phone_columns=["Phone"],
    category_columns=["Region"],
    group_by="Region",
    sum_column="Amount",
)


@pytest.fixture
def cleaned():
    return clean(read_table(SAMPLE), OPTS)


def test_headers_and_text_are_trimmed(cleaned):
    df, _ = cleaned
    assert list(df.columns)[:2] == ["Order ID", "Customer Name"]
    assert df.loc[0, "Customer Name"] == "An Nguyen"


def test_blank_and_duplicate_rows_removed(cleaned):
    df, report = cleaned
    assert report.blank_rows_removed == 1
    assert report.duplicates_removed == 1
    assert list(df["Order ID"]) == ["A-001", "A-002", "A-003", "A-004", "A-005"]


def test_categories_normalised(cleaned):
    df, _ = cleaned
    assert set(df["Region"]) == {"North", "South", "Central"}


def test_emails_lowercased_and_invalid_reported(cleaned):
    df, report = cleaned
    assert df.loc[0, "Email"] == "an.nguyen@example.com"
    assert {"Row": 3, "Column": "Email", "Value": "binh.tran@example", "Problem": "invalid email"} in report.issues


def test_phones_in_international_format(cleaned):
    df, report = cleaned
    assert list(df["Phone"][:4]) == ["+84901234567", "+84912345678", "+842838221234", "+84933111222"]
    assert any(i["Column"] == "Phone" and i["Value"] == "123" for i in report.issues)


@pytest.mark.parametrize(
    "raw, expected",
    [
        ("1,250,000", 1250000.0), ("$45.50", 45.5), ("1.250,75", 1250.75), ("12,5", 12.5),
        ("890000", 890000.0), ("45.500 ₫", 45500.0), ("1.250.000", 1250000.0), ("120.000đ", 120000.0),
        ("35.000 VND", 35000.0), ("-12.5", -12.5),
    ],
)
def test_number_formats(raw, expected):
    df, _ = clean(pd.DataFrame({"Amount": [raw]}), Options(number_columns=["Amount"]))
    assert df.loc[0, "Amount"] == expected


def test_bad_numbers_and_dates_are_reported_not_fatal(cleaned):
    df, report = cleaned
    problems = {(i["Row"], i["Problem"]) for i in report.issues}
    assert (7, "unrecognised date") in problems
    assert (8, "not a number") in problems
    assert pd.isna(df.loc[4, "Amount"])


def test_dates_read_day_first(cleaned):
    df, _ = cleaned
    assert df.loc[0, "Order Date"] == pd.Timestamp("2026-10-03")
    assert df.loc[1, "Order Date"] == pd.Timestamp("2026-10-04")
    assert df.loc[2, "Order Date"] == pd.Timestamp("2026-10-05")


def test_unknown_column_gives_clear_error():
    with pytest.raises(ValueError, match="Columns not found: Nope"):
        clean(read_table(SAMPLE), Options(date_columns=["Nope"]))


def test_report_has_three_sheets_and_chart(cleaned, tmp_path):
    df, report = cleaned
    out = tmp_path / "out.xlsx"
    write_report(df, report, summarize(df, OPTS), out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Cleaned", "Summary", "Issues"]
    assert len(wb["Summary"]._charts) == 1
    assert wb["Issues"].max_row == len(report.issues) + 1


def test_cli_end_to_end(tmp_path, capsys):
    out = tmp_path / "report.xlsx"
    code = main([str(SAMPLE), "-o", str(out), "--key", "Order ID", "--dates", "Order Date",
                 "--numbers", "Amount", "--emails", "Email", "--phones", "Phone",
                 "--categories", "Region", "--group-by", "Region", "--sum", "Amount"])
    assert code == 0
    assert out.exists()
    assert "7 rows in -> 5 rows out" in capsys.readouterr().out


@pytest.mark.parametrize(
    "raw, expected",
    [("2026/10/04", "2026-10-04"), ("20261004", "2026-10-04"), ("2026-10-04T10:00:00+07:00", "2026-10-04 10:00")],
)
def test_year_first_and_timezone_dates(raw, expected):
    df, report = clean(pd.DataFrame({"D": [raw]}), Options(date_columns=["D"]))
    assert df.loc[0, "D"] == pd.Timestamp(expected)
    assert df["D"].dt.tz is None
    assert not report.issues


def test_timezone_dates_can_be_written(tmp_path):
    df, report = clean(pd.DataFrame({"D": ["2026-10-04T10:00:00+07:00"]}), Options(date_columns=["D"]))
    write_report(df, report, None, tmp_path / "out.xlsx")


def test_accounting_negative():
    df, _ = clean(pd.DataFrame({"A": ["(1,234.50)"]}), Options(number_columns=["A"]))
    assert df.loc[0, "A"] == -1234.5


def test_issue_rows_match_file_after_blank_line(tmp_path):
    src = tmp_path / "in.csv"
    src.write_text("ID,Email\nA,bad1\n\nB,bad2\n", encoding="utf-8")
    _, report = clean(read_table(src), Options(email_columns=["Email"]))
    assert [i["Row"] for i in report.issues] == [2, 4]
    assert report.blank_rows_removed == 1


def test_semicolon_csv(tmp_path):
    src = tmp_path / "in.csv"
    src.write_text("ID;Amount\nA;1,5\n", encoding="utf-8")
    df, _ = clean(read_table(src), Options(number_columns=["Amount"]))
    assert df.loc[0, "Amount"] == 1.5


def test_sum_column_is_parsed_even_without_numbers_flag():
    opts = Options(group_by="G", sum_column="A")
    df, _ = clean(pd.DataFrame({"G": ["x", "x"], "A": ["1", "2"]}), opts)
    assert summarize(df, opts).loc[0, "Total A"] == 3


def test_rows_with_empty_key_are_not_deduplicated():
    df, report = clean(pd.DataFrame({"ID": [None, None, "A", "A"], "X": ["p", "q", "r", "r"]}), Options(key_columns=["ID"]))
    assert len(df) == 3
    assert report.duplicates_removed == 1


def test_formula_like_text_stays_text(tmp_path):
    df, report = clean(pd.DataFrame({"Note": ["=== TOTAL ===", '=HYPERLINK("http://x","y")']}), Options())
    out = tmp_path / "out.xlsx"
    write_report(df, report, None, out)
    ws = load_workbook(out)["Cleaned"]
    assert ws["A2"].data_type == "s" and ws["A2"].value == "=== TOTAL ==="
    assert ws["A3"].data_type == "s"


def test_phone_edge_cases():
    df, report = clean(pd.DataFrame({"P": ["+84 (0)912 345 678", "abc"]}), Options(phone_columns=["P"]))
    assert df.loc[0, "P"] == "+84912345678"
    assert df.loc[1, "P"] == "abc"
    assert any(i["Problem"] == "not a phone number" for i in report.issues)


def test_cli_reports_unwritable_output(tmp_path, capsys):
    code = main([str(SAMPLE), "-o", str(tmp_path / "missing-dir" / "out.xlsx")])
    assert code == 1
    assert "cannot write" in capsys.readouterr().err
