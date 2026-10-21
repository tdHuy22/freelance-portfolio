"""Runs the cleaning steps in order and writes a three-sheet Excel report."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import cleaning
from .cleaning import CleanReport


@dataclass
class Options:
    key_columns: list[str] = field(default_factory=list)
    date_columns: list[str] = field(default_factory=list)
    number_columns: list[str] = field(default_factory=list)
    email_columns: list[str] = field(default_factory=list)
    phone_columns: list[str] = field(default_factory=list)
    category_columns: list[str] = field(default_factory=list)
    group_by: str | None = None
    sum_column: str | None = None
    dayfirst: bool = True
    country_code: str = "84"


def read_table(path: Path, sheet: str | int = 0) -> pd.DataFrame:
    if path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
        return pd.read_excel(path, sheet_name=sheet, dtype=object)
    # sep=None sniffs commas vs semicolons; blank lines are kept so issue row numbers
    # match the file (drop_blank_rows removes and counts them).
    return pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig",
                       sep=None, engine="python", skip_blank_lines=False)


def _check_columns(df: pd.DataFrame, opts: Options) -> None:
    wanted = {
        *opts.key_columns, *opts.date_columns, *opts.number_columns, *opts.email_columns,
        *opts.phone_columns, *opts.category_columns,
        *(c for c in (opts.group_by, opts.sum_column) if c),
    }
    missing = sorted(wanted - set(df.columns))
    if missing:
        raise ValueError(f"Columns not found: {', '.join(missing)}. Available: {', '.join(df.columns)}")


def clean(df: pd.DataFrame, opts: Options) -> tuple[pd.DataFrame, CleanReport]:
    report = CleanReport(rows_in=len(df))
    df = cleaning.normalize_headers(df)
    _check_columns(df, opts)
    # Spreadsheet row numbers (header is row 1) make the issues sheet easy to act on.
    df.index = range(2, len(df) + 2)
    df = cleaning.strip_text(df)
    df = cleaning.drop_blank_rows(df, report)
    df = cleaning.title_case(df, opts.category_columns)
    df = cleaning.clean_emails(df, opts.email_columns, report)
    df = cleaning.clean_phones(df, opts.phone_columns, opts.country_code, report)
    # A summed column must be numeric, or pandas would concatenate the text.
    numbers = opts.number_columns + [c for c in [opts.sum_column] if c and c not in opts.number_columns]
    df = cleaning.parse_numbers(df, numbers, report)
    df = cleaning.parse_dates(df, opts.date_columns, opts.dayfirst, report)
    df = cleaning.drop_duplicates(df, opts.key_columns, report)
    report.rows_out = len(df)
    return df.reset_index(drop=True), report


def summarize(df: pd.DataFrame, opts: Options) -> pd.DataFrame | None:
    if not opts.group_by:
        return None
    grouped = df.groupby(opts.group_by, dropna=False)
    if opts.sum_column:
        table = grouped[opts.sum_column].agg(["count", "sum"]).rename(
            columns={"count": "Rows", "sum": f"Total {opts.sum_column}"}
        )
        table = table.sort_values(f"Total {opts.sum_column}", ascending=False)
    else:
        table = grouped.size().to_frame("Rows").sort_values("Rows", ascending=False)
    return table.reset_index()


HEADER_FILL = PatternFill("solid", fgColor="1F4E79")


def _style_header(ws, row: int) -> None:
    for cell in ws[row]:
        if cell.value is not None:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL


def _neutralise_formulas(ws) -> None:
    """Client data such as '=== TOTAL ===' or '=HYPERLINK(...)' stays text, never a live formula."""
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                cell.data_type = "s"


def _style_sheet(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"
    for idx, column in enumerate(ws.columns, start=1):
        width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 10), 50)


def write_report(df: pd.DataFrame, report: CleanReport, summary: pd.DataFrame | None, out: Path) -> None:
    overview = pd.DataFrame(
        [
            ("Rows in input", report.rows_in),
            ("Blank rows removed", report.blank_rows_removed),
            ("Duplicate rows removed", report.duplicates_removed),
            ("Rows in output", report.rows_out),
            ("Cells needing review", len(report.issues)),
        ],
        columns=["Step", "Count"],
    )
    issues = pd.DataFrame(report.issues, columns=["Row", "Column", "Value", "Problem"]).sort_values("Row", kind="stable")

    with pd.ExcelWriter(out, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        df.to_excel(writer, sheet_name="Cleaned", index=False)
        overview.to_excel(writer, sheet_name="Summary", index=False)
        if summary is not None:
            summary.to_excel(writer, sheet_name="Summary", index=False, startrow=len(overview) + 2)
        issues.to_excel(writer, sheet_name="Issues", index=False)

        for ws in writer.book.worksheets:
            _neutralise_formulas(ws)
            _style_sheet(ws)
        if summary is not None:
            _style_header(writer.book["Summary"], len(overview) + 3)

        if summary is not None and len(summary) and summary.shape[1] >= 2:
            ws = writer.book["Summary"]
            first = len(overview) + 3  # header row of the summary table (1-based)
            value_col = summary.shape[1]
            chart = BarChart()
            chart.title = summary.columns[-1]
            chart.legend = None
            chart.add_data(Reference(ws, min_col=value_col, min_row=first, max_row=first + len(summary)), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=1, min_row=first + 1, max_row=first + len(summary)))
            ws.add_chart(chart, "E2")
